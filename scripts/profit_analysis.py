#!/usr/bin/env python3
"""仟上私募客户盈利多维分析。

用法:
    python3 profit_analysis.py 目前客户持仓收益.xlsx [历史客户持仓收益.xlsx] [--asof 2026-06-30]

输入: 客户持仓收益xlsx，需含 `_总表` 开头的sheet，列头包含
    产品名称/客户名称/首次申/认购时间/累计本金(含费)/累计收益-扣除费用/累计收益率(扣除费用)
输出: 打印报告所需全部盈利指标（按人口径为主，按笔口径对照）。

口径（已拍板，见SKILL.md）:
- 盈利比例按人: 同一客户名下全部持仓的累计收益合并, >0 为盈利
- 持有时长: 该客户首次申购日 → asof
- 累计总盈利: 在持 + 历史(如提供) 逐笔加总
"""
import argparse
import re
import statistics
from datetime import date, datetime
from collections import defaultdict

import openpyxl
from openpyxl.utils.datetime import from_excel


def find_sheet(wb, prefix="_总表"):
    for s in wb.sheetnames:
        if s.startswith(prefix):
            return s
    raise SystemExit(f"未找到 {prefix} 开头的sheet, 现有: {wb.sheetnames}")


def col_index(header, kw):
    for i, c in enumerate(header):
        if c and kw in str(c):
            return i
    raise SystemExit(f"表头未找到含「{kw}」的列: {header}")


def normalize_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) < 8:
        return None
    try:
        return date(int(digits[:4]), int(digits[4:6]), int(digits[6:8]))
    except ValueError:
        return None


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[find_sheet(wb)]
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    ip, iname, ibuy = col_index(h, "产品名称"), col_index(h, "客户名称"), col_index(h, "首次申")
    iprin, iprof = col_index(h, "累计本金"), col_index(h, "累计收益-扣除")
    out = []
    for r in rows[1:]:
        prof = r[iprof]
        if r[iname] is None or not isinstance(prof, (int, float)):
            continue
        out.append(dict(prod=str(r[ip]), name=str(r[iname]).strip(),
                        buy=normalize_date(r[ibuy], wb.epoch),
                        principal=r[iprin] or 0, profit=prof))
    wb.close()
    return out


def yrs(b, asof):
    return (asof - b).days / 365.25 if isinstance(b, date) else None


def pct(a, b):
    return f"{a}/{b} = {a / b * 100:.1f}%" if b else "n/a"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("current", help="目前客户持仓收益 xlsx")
    parser.add_argument("historical", nargs="?", help="历史客户持仓收益 xlsx")
    parser.add_argument("--asof", type=date.fromisoformat, default=date.today(),
                        help="统计截止日，格式 YYYY-MM-DD；默认今天")
    args = parser.parse_args()
    asof = args.asof

    cur = load(args.current)
    his = load(args.historical) if args.historical else []

    print(f"数据截止: {asof}  在持{len(cur)}笔  历史{len(his)}笔\n")

    # ── 累计总盈利（全口径）
    tc, th = sum(x["profit"] for x in cur), sum(x["profit"] for x in his)
    print(f"◆ 累计总盈利(在持+历史全口径): {(tc + th) / 1e8:.4f}亿  (在持{tc / 1e8:.4f}亿 + 历史{th / 1e8:.4f}亿)")

    # ── 按人合并（在持）
    by = defaultdict(lambda: dict(profit=0.0, principal=0.0, first=None))
    for x in cur:
        p = by[x["name"]]
        p["profit"] += x["profit"]; p["principal"] += x["principal"]
        if x["buy"] and (p["first"] is None or x["buy"] < p["first"]):
            p["first"] = x["buy"]
    people = list(by.values())
    n = len(people)
    pos = [p for p in people if p["profit"] > 0]
    neg = [p for p in people if p["profit"] <= 0]
    print(f"\n◆ 在持按人口径: 投资者{n}人, 盈利 {pct(len(pos), n)}")
    kb = sum(1 for x in cur if x["profit"] > 0)
    print(f"  (对照·按笔口径: {pct(kb, len(cur))} — 报告文字写'投资者'时必须用按人口径)")

    # ── 持有时长分组
    print("\n◆ 按持有时长(首笔申购日起算):")
    for lab, lo, hi in [("4.5年以上", 4.5, 99), ("3-4.5年", 3, 4.5), ("1-3年", 1, 3), ("1年以内", 0, 1)]:
        sel = [p for p in people if p["first"] and lo <= (yrs(p["first"], asof) or -1) < hi]
        if not sel:
            continue
        w = sum(1 for p in sel if p["profit"] > 0)
        rts = [p["profit"] / p["principal"] for p in sel if p["principal"] > 0]
        med = f"{statistics.median(rts) * 100:.1f}%" if rts else "-"
        avg = sum(p["profit"] for p in sel) / len(sel)
        print(f"  {lab:<10} {len(sel):>4}人  盈利{pct(w, len(sel)):>16}  收益率中位数{med:>8}  人均盈亏{avg:>12,.0f}元")
    ge1 = [p for p in people if p["first"] and (yrs(p["first"], asof) or -1) >= 1]
    w1 = sum(1 for p in ge1 if p["profit"] > 0)
    print(f"  → 满1年以上合计: 盈利 {pct(w1, len(ge1))}")

    # ── 盈利金额分布
    wins = sorted((p["profit"] for p in pos), reverse=True)
    if wins:
        print(f"\n◆ 盈利金额(按人): 盈利{len(wins)}人")
        for th_ in (1e6, 5e5, 1e5):
            print(f"  盈利≥{th_ / 1e4:.0f}万: {sum(1 for w in wins if w >= th_)}人")
        print(f"  中位数{statistics.median(wins):,.0f}元  人均{statistics.mean(wins):,.0f}元")

    # ── 浮亏画像
    if neg:
        rts = [p["profit"] / p["principal"] for p in neg if p["principal"] > 0]
        prods = defaultdict(int)
        for x in cur:
            if by[x["name"]]["profit"] <= 0:
                prods[x["prod"][:14]] += 1
        print(f"\n◆ 浮亏画像: {len(neg)}人 (占比{len(neg) / n * 100:.1f}%)")
        if rts:
            print(f"  浮亏率中位数{statistics.median(rts) * 100:.2f}%  最深{min(rts) * 100:.2f}%  人均{sum(p['profit'] for p in neg) / len(neg):,.0f}元")
        else:
            print(f"  浮亏率无法计算（缺少有效本金）  人均{sum(p['profit'] for p in neg) / len(neg):,.0f}元")
        mx = max((yrs(p["first"], asof) or 0) for p in neg)
        print(f"  最长持有{mx:.2f}年  涉及产品: {dict(prods)}")

    # ── 分产品(按笔, 单产品内=按人)
    print("\n◆ 分产品(在持):")
    pp = defaultdict(lambda: defaultdict(float))
    for x in cur:
        pp[x["prod"]][x["name"]] += x["profit"]
    for prod, by_name in sorted(pp.items(), key=lambda kv: -len(kv[1])):
        vals = list(by_name.values())
        w = sum(1 for v in vals if v > 0)
        print(f"  {prod[:24]:<26} {len(vals):>3}人  盈利{pct(w, len(vals))}")


if __name__ == "__main__":
    main()
